"""Проверки исправлений по замечаниям ревью: свежесть, расчёт состояния школы, осторожность
ML, ТЗ (пороги, инциденты, линии, экспорт), сохранность истории, изоляция доступа.

Запуск:  cd backend && python -m unittest discover -s tests -v
Работает на временной SQLite-базе; боевую БД не трогает.
"""
import asyncio
import io
import os
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import AsyncMock, patch

_TMP = tempfile.mkdtemp()
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/test.db"
os.environ["EXTERNAL_POLL_ENABLED"] = "false"
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "agent"))

from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402

import app.database as database  # noqa: E402
from app.cache import cache_invalidate  # noqa: E402
from app.config import settings  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Device, Incident, Line, Measurement, School, SyncBatch, User  # noqa: E402
from app.security import create_device_token, fingerprint, hash_password  # noqa: E402
from app.services import smart_sync  # noqa: E402
from app.services.lines import ensure_defaults  # noqa: E402
from app.services.ml import attribution  # noqa: E402
from app.services.ml.features import Topology  # noqa: E402
from app.services.status import (  # noqa: E402
    STATUS_CRITICAL, STATUS_NO_DATA, STATUS_OFFLINE, STATUS_OK, STATUS_UNSTABLE, classify,
    freshness, reset_thresholds,
)

NOW = datetime.utcnow().replace(microsecond=0)
CLIENT = None   # TestClient, создаётся в setUpModule
TOKENS: dict[str, dict] = {}
HW = {"mac": "AA:BB", "cpu": "cpu", "os": "os"}    # оборудование для перерегистрации

# код школы → (провайдер, район)
SCHOOLS = {"A": ("P1", "R1"), "B": ("P1", "R1"), "B2": ("P1", "R1"), "C": ("P2", "R1"),
           "D": ("P3", "R2"), "E": ("P1", "R1"), "F": ("P4", "R3"), "G": ("P5", "R4")}


def gw(code): return f"PC-{code}-GW"
def ws(code): return f"PC-{code}-WS"


def setUpModule():
    global CLIENT
    # Если app уже импортирован с другой БД (общий запуск тестов) — переключаем сессии на временную.
    engine = create_engine(f"sqlite:///{_TMP}/test.db", connect_args={"check_same_thread": False})
    database.engine = engine
    database.SessionLocal.configure(bind=engine)
    database.migrate_schema()
    with database.SessionLocal() as db:
        for code, (provider, region) in SCHOOLS.items():
            school = School(school_id_code=code, name=f"Школа {code}", provider=provider, region=region,
                            connection_type="Оптика", contract_speed_down=100.0,
                            contract_speed_up=100.0, lat=50.0, lng=82.0)
            db.add(school)
            db.flush()
            for did, kind in ((gw(code), "Шлюз"), (ws(code), "Ноутбук")):
                db.add(Device(device_id=did, school_id=school.id, name=did, room="к1", device_type=kind,
                              hardware_fingerprint=fingerprint(HW["mac"], did, HW["cpu"], HW["os"]),
                              revoked=False))
        ids = {s.school_id_code: s.id for s in db.query(School)}
        for email, role, kw in (
                ("admin@t", "admin", {}), ("school@t", "school", {"school_id": ids["A"]}),
                ("p1@t", "provider", {"provider_name": "P1"}), ("p3@t", "provider", {"provider_name": "P3"}),
                ("dist@t", "district", {"district": "R1"}), ("lost@t", "school", {})):
            db.add(User(email=email, full_name=email, role=role, password_hash=hash_password("pw"), **kw))
        db.commit()
        ensure_defaults(db)
    CLIENT = TestClient(app)
    CLIENT.__enter__()


def tearDownModule():
    CLIENT.__exit__(None, None, None)


def auth(email: str) -> dict:
    if email not in TOKENS:
        token = CLIENT.post("/api/auth/login", json={"email": email, "password": "pw"}).json()["access_token"]
        TOKENS[email] = {"Authorization": f"Bearer {token}"}
    return TOKENS[email]


def m(minutes_ago: float, down=95.0, up=90.0, ping=20.0, jitter=2.0, loss=0.0, offline=False) -> dict:
    return {"timestamp": (NOW - timedelta(minutes=minutes_ago)).isoformat(), "download_speed": down,
            "upload_speed": up, "ping": ping, "jitter": jitter, "packet_loss": loss,
            "is_offline": offline, "source": "live"}


BAD = dict(down=0.5, up=0.2, ping=400.0, jitter=90.0, loss=30.0)
CRIT = dict(down=5.0, up=90.0, ping=20.0, jitter=2.0, loss=0.0)   # скорость 5% договора


def ingest(device_id: str, *items: dict) -> int:
    with database.SessionLocal() as db:
        device = db.query(Device).filter(Device.device_id == device_id).first()
        stored = smart_sync.ingest(db, device, list(items))
        db.commit()
    cache_invalidate("dash:")
    return stored


def school(code: str) -> School:
    with database.SessionLocal() as db:
        s = db.query(School).filter(School.school_id_code == code).one()
        db.expunge(s)
        return s


def school_id(code: str) -> int:
    return school(code).id


class Thresholds(unittest.TestCase):
    def test_tz_thresholds_and_all_five_metrics_affect_status(self):
        self.assertEqual(classify(95, 20, 0, 100, upload=90, jitter=2, contract_up=100), STATUS_OK)
        self.assertEqual(classify(95, 101, 0, 100), STATUS_UNSTABLE)                     # Ping > 100
        self.assertEqual(classify(95, 20, 0, 100, upload=19, contract_up=100), STATUS_UNSTABLE)  # Upload < 20
        self.assertEqual(classify(95, 20, 0, 100, jitter=31), STATUS_UNSTABLE)           # Jitter > 30
        self.assertEqual(classify(95, 20, 2.5, 100), STATUS_UNSTABLE)                    # Loss > 2
        self.assertEqual(classify(95, 20, 6, 100), STATUS_CRITICAL)
        self.assertEqual(classify(0, 0, 100, 100), STATUS_OFFLINE)
        # договор ниже базового порога 20 Мбит/с судится по договору, а не «нарушает всегда»
        self.assertEqual(classify(7, 20, 0, 10, upload=30, contract_up=10), STATUS_OK)

    def test_admin_changes_thresholds_and_measurement_keeps_applied_version(self):
        body = {"down_min": 20, "up_min": 20, "ping_max": 50, "jitter_max": 30, "loss_max": 2,
                "availability_min": 99, "contract_ratio": 0.6, "incident_after": 3, "stale_after_min": 90}
        self.assertEqual(CLIENT.put("/api/admin/thresholds", json=body, headers=auth("p1@t")).status_code, 403)
        r = CLIENT.put("/api/admin/thresholds", json=body, headers=auth("admin@t"))
        try:
            self.assertEqual(r.status_code, 200)
            version = r.json()["active"]["id"]
            ingest(gw("B"), m(3, ping=60))                       # 60 мс: при пороге 100 была бы норма
            with database.SessionLocal() as db:
                row = db.query(Measurement).filter(Measurement.device_id == gw("B")).order_by(
                    Measurement.timestamp.desc()).first()
                self.assertEqual((row.status, row.threshold_id), (STATUS_UNSTABLE, version))
        finally:
            CLIENT.put("/api/admin/thresholds", json={**body, "ping_max": 100}, headers=auth("admin@t"))
            reset_thresholds()


class Freshness(unittest.TestCase):
    def test_old_data_is_no_data_not_last_known_status(self):
        with database.SessionLocal() as db:
            s = db.query(School).filter(School.school_id_code == "F").one()
            s.status, s.last_measurement = STATUS_OK, NOW - timedelta(days=4)
            db.commit()
        cache_invalidate("dash:")
        rows = {r["school_id_code"]: r for r in CLIENT.get("/api/web/schools", headers=auth("admin@t")).json()}
        f = rows["F"]
        self.assertEqual((f["status"], f["last_known_status"], f["is_stale"]), (STATUS_NO_DATA, STATUS_OK, True))
        self.assertIsNotNone(f["last_measurement"])
        # фильтр карты по статусу: F не «Норма», а «Нет свежих данных»
        ok = CLIENT.get("/api/web/schools", params={"status": STATUS_OK}, headers=auth("admin@t")).json()
        self.assertNotIn("F", {r["school_id_code"] for r in ok})
        # режим демонстрации истории: то же состояние «на момент» последнего замера
        at = (NOW - timedelta(days=4)).isoformat()
        demo = {r["school_id_code"]: r for r in CLIENT.get(
            "/api/web/schools", params={"as_of": at}, headers=auth("admin@t")).json()}
        self.assertEqual(demo["F"]["status"], STATUS_OK)

    def test_overview_reports_freshness_and_history_mode(self):
        old = (NOW - timedelta(days=30)).isoformat()
        ov = CLIENT.get("/api/web/overview", params={"as_of": old}, headers=auth("admin@t")).json()
        self.assertEqual(ov["freshness"]["mode"], "history")
        self.assertTrue(ov["freshness"]["is_stale"])
        self.assertIsNone(ov["sla_compliance"])                  # нет данных — нет процента «в норме»
        live = CLIENT.get("/api/web/overview", headers=auth("admin@t")).json()
        self.assertEqual(live["freshness"]["mode"], "live")
        self.assertIn("no_data", live["status_counts"])

    def test_freshness_helper(self):
        self.assertTrue(freshness(None, NOW)["is_stale"])
        self.assertFalse(freshness(NOW - timedelta(minutes=10), NOW)["is_stale"])
        self.assertEqual(freshness(NOW - timedelta(hours=3), NOW)["age_min"], 180)

    def test_empty_ml_lists_come_with_freshness(self):
        summary = CLIENT.get("/api/ml/summary", params={"as_of": (NOW - timedelta(days=30)).isoformat()},
                             headers=auth("admin@t")).json()
        self.assertTrue(summary["freshness"]["is_stale"])
        self.assertEqual(CLIENT.get("/api/ml/board", params={"as_of": (NOW - timedelta(days=30)).isoformat()},
                                    headers=auth("admin@t")).json(), [])


class SchoolState(unittest.TestCase):
    def test_workstation_and_old_backfill_do_not_change_school(self):
        ingest(gw("G"), m(20))                                    # основная линия в норме
        s = school("G")
        self.assertEqual((s.status, s.current_download), (STATUS_OK, 95.0))
        ingest(ws("G"), m(5, **BAD))                              # плохой Wi-Fi ноутбука
        s = school("G")
        self.assertEqual((s.status, s.current_download), (STATUS_OK, 95.0))
        ingest(gw("G"), m(120, **CRIT))                           # старая офлайн-догрузка
        s = school("G")
        self.assertEqual((s.status, s.current_download), (STATUS_OK, 95.0))
        self.assertEqual(s.last_measurement, NOW - timedelta(minutes=20))
        with database.SessionLocal() as db:                       # история при этом сохранена
            self.assertEqual(db.query(Measurement).filter(Measurement.device_id == gw("G")).count(), 2)
        detail = CLIENT.get(f"/api/web/schools/{s.id}", headers=auth("admin@t")).json()
        self.assertEqual(detail["workstations"]["degraded"], 1)
        self.assertEqual(detail["workstations"]["total"], 1)
        self.assertEqual([(x["code"], x["role"]) for x in detail["lines"]], [("G-L1", "main")])
        ingest(gw("G"), m(1, **CRIT))                             # а вот новый замер линии учитывается
        self.assertEqual(school("G").status, STATUS_CRITICAL)

    def test_backup_line_is_tracked_separately(self):
        with database.SessionLocal() as db:
            sid = db.query(School).filter(School.school_id_code == "B2").one().id
            backup = Line(school_id=sid, code="B2-L2", role="backup", provider="Q", contract_speed_down=50)
            db.add(backup)
            db.flush()
            db.query(Device).filter(Device.device_id == ws("B2")).update({"line_id": backup.id})
            db.commit()
        ingest(gw("B2"), m(30))
        ingest(ws("B2"), m(2, **CRIT))                            # точка мониторинга резервной линии
        s = school("B2")
        self.assertEqual(s.status, STATUS_OK)                     # школу резервная линия не красит
        detail = CLIENT.get(f"/api/web/schools/{s.id}", headers=auth("admin@t")).json()
        by_code = {x["code"]: x for x in detail["lines"]}
        self.assertEqual((by_code["B2-L1"]["role"], by_code["B2-L1"]["status"]), ("main", STATUS_OK))
        self.assertEqual((by_code["B2-L2"]["role"], by_code["B2-L2"]["status"]), ("backup", STATUS_CRITICAL))

    def test_incident_needs_consecutive_bad_measurements(self):
        def incidents():
            with database.SessionLocal() as db:
                return db.query(Incident).filter(Incident.school_id == school_id("E")).all()

        ingest(gw("E"), m(60, **CRIT), m(50, **CRIT))             # два плохих — ещё не инцидент
        self.assertEqual(incidents(), [])
        ingest(gw("E"), m(30))                                    # норма сбрасывает серию
        ingest(gw("E"), m(20, **CRIT), m(10, **CRIT))
        self.assertEqual(incidents(), [])
        ingest(gw("E"), m(1, **CRIT))                             # третий подряд
        found = incidents()
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0].start_time, NOW - timedelta(minutes=20))
        ingest(gw("E"), m(0.5, offline=True, down=0, up=0, ping=0, jitter=0, loss=100))
        self.assertEqual(len(incidents()), 1)                     # открытый инцидент не дублируется


class Durability(unittest.TestCase):
    def _device_headers(self, code):
        with database.SessionLocal() as db:
            d = db.query(Device).filter(Device.device_id == gw(code)).one()
            token = create_device_token(d.device_id, d.school_id, d.hardware_fingerprint)
        return {"Authorization": f"Bearer {token}"}

    def test_accepted_batch_is_in_db_before_processing_and_survives_lost_queue(self):
        items = [m(90 + i * 10) for i in range(3)]
        with patch("app.routers.agent.enqueue", new=AsyncMock(return_value=True)):    # «процесс упал» до воркера
            r = CLIENT.post("/api/agent/measurements/batch", json={"items": items}, headers=self._device_headers("B"))
        self.assertEqual(r.status_code, 202)
        batch_id = r.json()["batch_id"]
        with database.SessionLocal() as db:
            batch = db.get(SyncBatch, batch_id)
            self.assertEqual(batch.status, "queued")
            self.assertTrue(batch.payload)                                             # замеры уже в БД
            before = db.query(Measurement).filter(Measurement.device_id == gw("B")).count()
        smart_sync.persist_batch(batch_id)                                             # то, что делает восстановление
        with database.SessionLocal() as db:
            batch = db.get(SyncBatch, batch_id)
            self.assertEqual((batch.status, batch.payload), ("done", None))
            after = db.query(Measurement).filter(Measurement.device_id == gw("B")).count()
        self.assertEqual(after - before, 3)
        smart_sync.persist_batch(batch_id)                                             # повтор — без последствий
        # агент повторно прислал тот же пакет (не получил ответ) — дублей нет
        with patch("app.routers.agent.enqueue", new=AsyncMock(return_value=True)):
            again = CLIENT.post("/api/agent/measurements/batch", json={"items": items},
                                headers=self._device_headers("B")).json()["batch_id"]
        smart_sync.persist_batch(again)
        with database.SessionLocal() as db:
            self.assertEqual(db.query(Measurement).filter(Measurement.device_id == gw("B")).count(), after)
        state = CLIENT.get(f"/api/agent/measurements/batch/{batch_id}", headers=self._device_headers("B")).json()
        self.assertEqual(state["status"], "done")
        # чужой пакет недоступен
        self.assertEqual(CLIENT.get(f"/api/agent/measurements/batch/{batch_id}",
                                    headers=self._device_headers("A")).status_code, 404)

    def test_failed_batch_keeps_payload_and_recovery_picks_only_unfinished(self):
        with database.SessionLocal() as db:
            db.add_all([
                SyncBatch(device_id="нет-такого", school_id=1, items=1, status="queued", payload="[{}]"),
                SyncBatch(device_id=gw("B"), school_id=1, items=1, status="failed", payload="[]", attempts=1),
                SyncBatch(device_id=gw("B"), school_id=1, items=1, status="failed", payload="[]",
                          attempts=smart_sync.MAX_ATTEMPTS),
                SyncBatch(device_id=gw("B"), school_id=1, items=1, status="done", payload=None),
            ])
            db.commit()
            ghost = db.query(SyncBatch).filter(SyncBatch.device_id == "нет-такого").one().id
            wanted = {b.id for b in db.query(SyncBatch).filter(
                SyncBatch.status.in_(("queued", "processing", "failed")), SyncBatch.payload.isnot(None),
                SyncBatch.attempts < smart_sync.MAX_ATTEMPTS)}
        with self.assertRaises(LookupError):
            smart_sync.persist_batch(ghost)
        with database.SessionLocal() as db:
            b = db.get(SyncBatch, ghost)
            self.assertEqual((b.status, b.attempts, bool(b.payload)), ("failed", 1, True))   # ничего не потеряно
        queued = AsyncMock(return_value=True)
        with patch("app.services.smart_sync.enqueue", new=queued):
            asyncio.run(smart_sync.recover_pending())
        self.assertEqual({c.args[0] for c in queued.await_args_list} & wanted, wanted)
        with database.SessionLocal() as db:
            exhausted = {b.id for b in db.query(SyncBatch).filter(SyncBatch.attempts >= smart_sync.MAX_ATTEMPTS)}
        self.assertFalse({c.args[0] for c in queued.await_args_list} & exhausted)

    def test_agent_deletes_spool_only_after_server_confirms_write(self):
        os.environ["VKO_STATE_DIR"] = tempfile.mkdtemp()
        import agent

        class Reply:
            def __init__(self, code, body=None): self.status_code, self._body = code, body or {}
            def json(self): return self._body

        class Session:
            def __init__(self, reply): self.reply = reply
            def get(self, *a, **k): return self.reply

        conn = agent.spool_init()
        for i in range(3):
            agent.spool_add(conn, {"n": i})
        ids, _ = agent.spool_take(conn, 10)
        agent.spool_mark_sent(conn, ids, 7)                       # сервер ответил 202 — но это не «записано»
        self.assertEqual((agent.spool_size(conn), agent.spool_size(conn, unsent_only=True)), (3, 0))
        agent.confirm_sent(Session(Reply(200, {"status": "queued"})), conn)
        agent.confirm_sent(Session(Reply(200, {"status": "failed", "error": "x"})), conn)
        self.assertEqual(agent.spool_size(conn), 3)               # пока не done — копии на месте
        agent.confirm_sent(Session(Reply(404)), conn)             # сервер пакета не знает — отправить заново
        self.assertEqual(agent.spool_size(conn, unsent_only=True), 3)
        agent.spool_mark_sent(conn, ids, 8)
        agent.confirm_sent(Session(Reply(200, {"status": "done"})), conn)
        self.assertEqual(agent.spool_size(conn), 0)


class Isolation(unittest.TestCase):
    def test_trend_is_limited_to_users_school_and_main_line(self):
        ingest(gw("A"), m(15))
        ingest(gw("B"), m(14), m(13))
        ingest(ws("B"), m(12, **BAD))                             # рабочее место в тренд линий не входит
        cache_invalidate("dash:")
        own = CLIENT.get("/api/web/trend", headers=auth("school@t")).json()
        everyone = CLIENT.get("/api/web/trend", headers=auth("admin@t")).json()
        with database.SessionLocal() as db:
            expected = db.query(Measurement).filter(
                Measurement.school_id == school_id("A"), Measurement.device_id == gw("A"),
                Measurement.timestamp >= datetime.utcnow() - timedelta(hours=24)).count()
        self.assertEqual(sum(x["samples"] for x in own), expected)
        self.assertGreater(sum(x["samples"] for x in everyone), sum(x["samples"] for x in own))
        self.assertLess(max(x["ping"] for x in everyone), 100)    # ping рабочего места (400) не попал

    def test_caches_are_not_shared_between_scopes(self):
        p1 = CLIENT.get("/api/web/overview", headers=auth("p1@t")).json()["total_schools"]
        p3 = CLIENT.get("/api/web/overview", headers=auth("p3@t")).json()["total_schools"]
        dist = CLIENT.get("/api/web/overview", headers=auth("dist@t")).json()["total_schools"]
        self.assertEqual((p1, p3, dist), (4, 1, 5))
        self.assertEqual(len(CLIENT.get("/api/web/schools", headers=auth("dist@t")).json()), 5)

    def test_role_without_scope_sees_nothing(self):
        lost = auth("lost@t")
        self.assertEqual(CLIENT.get("/api/web/schools", headers=lost).json(), [])
        self.assertEqual(CLIENT.get("/api/web/overview", headers=lost).json()["total_devices"], 0)
        self.assertEqual(CLIENT.get(f"/api/web/schools/{school_id('A')}", headers=lost).status_code, 403)
        self.assertEqual(CLIENT.get("/api/ml/board", headers=lost).json(), [])
        self.assertEqual(CLIENT.get("/api/web/incidents", headers=lost).json(), [])

    def test_system_panel_is_not_for_school_role(self):
        self.assertEqual(CLIENT.get("/api/web/system", headers=auth("school@t")).status_code, 403)
        self.assertEqual(CLIENT.get("/api/web/system", headers=auth("admin@t")).status_code, 200)

    def test_revoked_device_cannot_reenroll_and_hardware_change_needs_admin(self):
        secret = fingerprint("enroll", "C", settings.JWT_SECRET)[:12]

        def enroll(**over):
            body = {"device_id": ws("C"), "school_id_code": "C", "enrollment_secret": secret,
                    "mac_address": HW["mac"], "cpu_model": HW["cpu"], "os_name": HW["os"], **over}
            return CLIENT.post("/api/agent/enroll", json=body).status_code

        self.assertEqual(enroll(), 200)                                   # то же оборудование — можно
        self.assertEqual(enroll(mac_address="EVIL"), 409)                 # перенос на другое — нельзя
        self.assertEqual(CLIENT.post(f"/api/admin/devices/{ws('C')}/revoke", headers=auth("admin@t")).status_code, 200)
        self.assertEqual(enroll(), 403)                                   # отозванное само не вернётся
        self.assertEqual(CLIENT.post(f"/api/admin/devices/{ws('C')}/reset-enrollment",
                                     headers=auth("admin@t")).status_code, 200)
        self.assertEqual(enroll(mac_address="NEW"), 200)                  # после сброса админом — можно


class Attribution(unittest.TestCase):
    def _snap(self, db, anomalous: set[str], present: set[str]) -> dict:
        state = lambda sid, bad: {"school_id": sid, "z": -8.0 if bad else 0.0, "depth": 0.6 if bad else 0.0,  # noqa: E731
                                  "anomaly": bad, "offline": False, "expected": 1.0,
                                  "ratio": 0.4 if bad else 1.0, "samples": 1}
        snap = {}
        for code in present:
            for did in (gw(code), ws(code)):
                snap[did] = state(school_id(code), code in anomalous)
        return snap

    def test_school_without_comparable_peers_gets_no_confident_verdict(self):
        with database.SessionLocal() as db:
            topo = Topology(db)
            snap = self._snap(db, anomalous={"D"}, present={"D"})     # у провайдера P3 в районе R2 соседей нет
            v = attribution.diagnose(db, school_id("D"), at=NOW, topo=topo, snap=snap)
        self.assertEqual(v["cause"], "undetermined")
        self.assertEqual(v["data_quality"]["level"], 0)
        self.assertLessEqual(v["confidence"], 0.3)
        self.assertIn("нет сопоставимых школ", v["narrative"])
        self.assertFalse(v["actionable"])
        self.assertIsNotNone(v["model_hint"])                         # гипотеза модели видна, но не выдаётся за вывод

    def test_peers_without_data_are_not_counted_as_healthy(self):
        with database.SessionLocal() as db:
            topo = Topology(db)
            snap = self._snap(db, anomalous={"A"}, present={"A"})     # B, B2 есть в справочнике, замеров нет
            _, n = topo._peer_fraction(topo.by_pair[("P1", "R1")], snap, school_id("A"))
        self.assertEqual(n, 0)

    def test_whole_school_verdict_is_hedged_and_not_grounds_for_claim(self):
        with database.SessionLocal() as db:
            topo = Topology(db)
            snap = self._snap(db, anomalous={"A"}, present={"A", "B", "B2", "E", "C"})
            v = attribution.diagnose(db, school_id("A"), at=NOW, topo=topo, snap=snap)
        self.assertEqual(v["cause"], "school_lan")
        self.assertTrue(v["narrative"].startswith("Предполагаемый источник"))
        self.assertIn("индивидуальная линия провайдера", v["narrative"])
        self.assertIn("не разделяет", v["narrative"])
        self.assertFalse(v["actionable"])
        self.assertLessEqual(v["confidence"], 0.9)

    def test_no_recent_measurements_is_no_data_not_normal(self):
        with database.SessionLocal() as db:
            v = attribution.diagnose(db, school_id("F"), at=NOW, snap={})
        self.assertEqual(v["cause"], "no_data")
        self.assertIn("Нет свежих данных", v["narrative"])
        self.assertNotEqual(v["cause"], "none")


class Export(unittest.TestCase):
    def test_csv_and_xlsx_with_filters_and_scope(self):
        ingest(gw("A"), m(200), m(190, **CRIT))
        sid = school_id("A")
        r = CLIENT.get("/api/web/export", params={"school_id": sid}, headers=auth("admin@t"))
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.headers["content-type"].startswith("text/csv"))
        lines = r.text.lstrip("﻿").splitlines()
        for column in ("Школа", "School ID", "Device ID", "Кабинет", "Дата (UTC)", "Время (UTC)",
                       "Download, Мбит/с", "Upload, Мбит/с", "Ping, мс", "Jitter, мс", "Packet Loss, %",
                       "Статус соединения"):
            self.assertIn(column, lines[0])
        with database.SessionLocal() as db:
            total = db.query(Measurement).filter(Measurement.school_id == sid).count()
        self.assertEqual(len(lines) - 1, total)
        bad = CLIENT.get("/api/web/export", params={"school_id": sid, "status": STATUS_CRITICAL},
                         headers=auth("admin@t")).text.lstrip("﻿").splitlines()
        self.assertTrue(all("Критично" in row for row in bad[1:]))
        self.assertGreaterEqual(len(bad), 2)
        one = CLIENT.get("/api/web/export", params={"school_id": sid, "device_id": ws("A")},
                         headers=auth("admin@t")).text.lstrip("﻿").splitlines()
        self.assertTrue(all(ws("A") in row for row in one[1:]))

        x = CLIENT.get("/api/web/export", params={"format": "xlsx", "school_id": sid}, headers=auth("admin@t"))
        self.assertEqual(x.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(x.content)) as z:
            self.assertIn("xl/worksheets/sheet1.xml", z.namelist())
            self.assertIn("Download", z.read("xl/worksheets/sheet1.xml").decode())
        try:
            import openpyxl
        except ImportError:
            pass
        else:
            sheet = openpyxl.load_workbook(io.BytesIO(x.content)).active
            self.assertEqual(sheet.max_row - 1, total)
            self.assertEqual(sheet.cell(1, 1).value, "Школа")

        summary = CLIENT.get("/api/web/export", params={"aggregate": "true", "school_id": sid},
                             headers=auth("admin@t")).text.lstrip("﻿").splitlines()
        self.assertIn("Доля проблемных, %", summary[0])
        self.assertEqual(len(summary), 2)

    def test_export_respects_role_scope_and_validates_fields(self):
        sid_d = school_id("D")
        self.assertEqual(CLIENT.get("/api/web/export", params={"school_id": sid_d},
                                    headers=auth("school@t")).status_code, 403)
        own = CLIENT.get("/api/web/export", headers=auth("school@t")).text.lstrip("﻿").splitlines()
        self.assertTrue(all("Школа A" in row for row in own[1:]))     # школьная роль видит только свою школу
        self.assertEqual(CLIENT.get("/api/web/export", params={"fields": "school,secret"},
                                    headers=auth("admin@t")).status_code, 400)


if __name__ == "__main__":
    unittest.main()
